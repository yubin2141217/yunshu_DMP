#!/usr/bin/env python3
"""
操作手册 Markdown → Word（本地嵌入图片）

硬性要求（见 references/word-export.md）：
1. 图片必须以二进制写入 docx（word/media/），禁止仅保留路径引用
2. 目录必须可点击跳转到对应章节（书签 + 内部超链接）
3. 操作步骤序号按章节/小节列表重新从 1 开始，禁止全文连续编号
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")
LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def set_run_font(run, name: str = "微软雅黑", size_pt: float | None = None, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.bold = bold


def add_inline_runs(paragraph, text: str, base_size: float = 11):
    text = LINK_RE.sub(r"\1", text)
    parts = BOLD_RE.split(text)
    for i, part in enumerate(parts):
        is_bold = i % 2 == 1
        code_parts = CODE_RE.split(part)
        for j, seg in enumerate(code_parts):
            if not seg:
                continue
            run = paragraph.add_run(seg)
            set_run_font(run, size_pt=base_size, bold=is_bold or (j % 2 == 1))
            if j % 2 == 1:
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def strip_front_matter(text: str) -> str:
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            return text[end + 4 :].lstrip("\n")
    return text


def is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def parse_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def bookmark_name_for(title: str, index: int) -> str:
    # Word 书签名：字母数字下划线，以字母或 _ 开头
    return f"sec_{index}"


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    """在段落开头插入书签起止标记，供目录超链接定位。"""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, text: str, bookmark_name: str) -> None:
    """目录项：点击跳转到书签（不依赖 TOC 域刷新）。"""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)

    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "微软雅黑")
    r_fonts.set(qn("w:hAnsi"), "微软雅黑")
    r_fonts.set(qn("w:eastAsia"), "微软雅黑")
    r_pr.append(r_fonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")  # 11pt
    r_pr.append(sz)

    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading_with_bookmark(
    doc: Document, title: str, level: int, bookmark_name: str, bookmark_id: int
):
    p = doc.add_heading(title, level=level)
    size = 16 if level == 1 else 14
    for r in p.runs:
        set_run_font(r, size_pt=size, bold=True)
    add_bookmark(p, bookmark_name, bookmark_id)
    return p


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[ci] if ci < len(row) else ""
            add_inline_runs(p, val, base_size=10)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def add_image(doc: Document, md_dir: Path, alt: str, rel_path: str) -> None:
    path = (md_dir / rel_path).resolve()
    if not path.is_file():
        p = doc.add_paragraph(f"[图片缺失: {rel_path}]")
        if p.runs:
            set_run_font(p.runs[0], size_pt=10)
        return
    width = Cm(15.5)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(path), width=width)
    if alt:
        cap = doc.add_paragraph(f"图：{alt}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            set_run_font(r, size_pt=9)
            r.italic = True
    print(f"  embedded: {rel_path} ({path.stat().st_size} bytes)")


def collect_h1_titles(lines: list[str]) -> list[str]:
    """收集一级章节标题（跳过「目录」本身），用于生成可跳转目录。"""
    titles: list[str] = []
    for line in lines:
        s = line.strip()
        if s.startswith("## ") and not s.startswith("### "):
            title = s[3:].strip()
            if title != "目录":
                titles.append(title)
    return titles


def add_numbered_step(doc: Document, number: str, text: str) -> None:
    """
    使用 Markdown 中的显式序号写入正文，避免 Word「列表编号」样式全文连续递增。
    每个章节在 MD 里从 1 重排时，Word 中也会从 1 开始。
    """
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    run = p.add_run(f"{number}. ")
    set_run_font(run, size_pt=11, bold=False)
    add_inline_runs(p, text, base_size=11)


def convert(md_file: Path, out_file: Path | None = None) -> Path:
    md_file = md_file.resolve()
    out_file = out_file or md_file.with_suffix(".docx")
    text = strip_front_matter(md_file.read_text(encoding="utf-8"))
    lines = text.splitlines()
    md_dir = md_file.parent

    h1_titles = collect_h1_titles(lines)
    # title -> bookmark
    h1_bookmarks: dict[str, tuple[str, int]] = {}
    for idx, title in enumerate(h1_titles, start=1):
        h1_bookmarks[title] = (bookmark_name_for(title, idx), idx)

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)

    i = 0
    embedded = 0
    next_bookmark_id = len(h1_titles) + 1
    h2_counter = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        # 目录：可点击跳转（内部超链接 → 章节书签）
        if stripped == "## 目录":
            p = doc.add_heading("目录", level=1)
            for r in p.runs:
                set_run_font(r, size_pt=16, bold=True)
            i += 1
            # 优先用正文实际 H1；若 MD 目录列表存在则按列表文案显示
            toc_items: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("- ["):
                toc_items.append(LINK_RE.sub(r"\1", lines[i].strip()[2:]))
                i += 1
            if not toc_items:
                toc_items = list(h1_titles)

            for item in toc_items:
                # 匹配书签：完整标题或「一、xxx」与 H1 相同
                bookmark = None
                for title, (bname, _) in h1_bookmarks.items():
                    if item == title or item in title or title in item:
                        bookmark = bname
                        break
                bp = doc.add_paragraph()
                bp.paragraph_format.space_after = Pt(4)
                if bookmark:
                    add_internal_hyperlink(bp, item, bookmark)
                else:
                    run = bp.add_run(item)
                    set_run_font(run, size_pt=11)
            continue

        img = IMAGE_RE.fullmatch(stripped)
        if img:
            add_image(doc, md_dir, img.group(1), img.group(2).strip())
            embedded += 1
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [parse_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            set_run_font(run, name="Consolas", size_pt=10)
            continue

        if stripped.startswith("### "):
            title = stripped[4:].strip()
            h2_counter += 1
            bname = f"sub_{h2_counter}"
            add_heading_with_bookmark(doc, title, 2, bname, next_bookmark_id)
            next_bookmark_id += 1
            i += 1
            continue

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if title == "目录":
                i += 1
                continue
            if title in h1_bookmarks:
                bname, bid = h1_bookmarks[title]
                add_heading_with_bookmark(doc, title, 1, bname, bid)
            else:
                add_heading_with_bookmark(
                    doc, title, 1, f"sec_extra_{next_bookmark_id}", next_bookmark_id
                )
                next_bookmark_id += 1
            i += 1
            continue

        if stripped.startswith("> "):
            quote = stripped[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline_runs(p, quote, base_size=10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        m_num = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_num:
            # 关键：不用 List Number 样式，避免跨章节连续编号
            add_numbered_step(doc, m_num.group(1), m_num.group(2))
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1

    doc.save(str(out_file))
    size = out_file.stat().st_size
    print(f"Saved: {out_file}")
    print(f"Embedded images: {embedded}")
    print(f"TOC hyperlinks: {len(h1_titles)}")
    print(f"File size: {size} bytes ({size / 1024:.1f} KB)")
    return out_file


def main() -> int:
    if len(sys.argv) < 2:
        print(
            "Usage: md_to_docx_embed.py <markdown> [output.docx]",
            file=sys.stderr,
        )
        return 1
    md = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    convert(md, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
